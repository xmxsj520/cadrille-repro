"""生成"Table 1 复现对比表"Word 文档：完整复刻论文表格结构，复现结果标黄，下方补 CAD-Recode 数据"""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

DST = r"E:\论文\Table1_复现对比.docx"

CN = "微软雅黑"
BODY = "宋体"

# (方法, 训练数据, PC_CD, PC_IoU, PC_IR, IMG_CD, IMG_IoU, IMG_IR, TXT_CD, TXT_IoU, TXT_IR,
#  加粗集合, 下划线集合)
PAPER = [
    ("PointNet→DeepCAD", "D_p", "9.64", "46.7", "7.1", "", "", "", "", "", "", set(), set()),
    ("Point-BERT→HNC-CAD", "D_p", "8.64", "65.3", "5.6", "", "", "", "", "", "", set(), set()),
    ("MultiCAD", "D_p", "8.09", "", "11.5", "", "", "", "", "", "", set(), set()),
    ("TransCAD", "D_p", "4.51", "65.5", "1.1", "", "", "", "", "", "", set(), set()),
    ("PrismCAD", "D_p", "4.28", "72.1", "16.2", "", "", "", "", "", "", set(), set()),
    ("Point2Cyl", "D_p", "4.27", "73.8", "3.9", "", "", "", "", "", "", set(), set()),
    ("CAD-Diffuser", "D_p", "3.02", "74.3", "1.5", "", "", "", "", "", "", set(), set()),
    ("CAD-SIGNet", "D_p", "0.29", "77.3", "5.0", "", "", "", "", "", "", {2, 3}, set()),
    ("DINOv2→HNC-CAD", "D_i", "", "", "", "2.08", "", "10.1", "", "", "", set(), set()),
    ("DINOv2→DeepCAD", "D_i", "", "", "", "1.13", "", "10.6", "", "", "", set(), set()),
    ("CADCrafter", "D_i", "", "", "", "0.26", "", "3.6", "", "", "", {5}, set()),
    ("BERT→DeepCAD", "D_t", "", "", "", "", "", "", "32.82", "", "10.0", set(), set()),
    ("CADmium", "D_t", "", "", "", "", "", "", "0.38", "", "4.3", set(), set()),
    ("Text2CAD", "D_t", "", "", "", "", "", "", "0.37", "71.5", "3.7", {9}, {9}),
    ("CAD-Coder", "D_t", "", "", "", "", "", "", "0.33", "", "5.3", set(), set()),
    ("Text-to-CadQuery", "D_t", "", "", "", "", "", "", "0.22", "", "1.3", {8}, {10}),
    ("cadrille", "D_pit", "0.25", "79.4", "0.4", "0.25", "78.2", "0.5", "0.21", "81.1", "1.4",
     {2, 3, 4, 5, 6, 7, 8, 9, 10}, {10}),
    ("CAD-Recode", "R_p", "0.18", "87.1", "3.1", "", "", "", "", "", "", set(), {3}),
    ("cadrille", "R_pi", "0.18", "87.1", "2.1", "0.18", "86.1", "1.5", "", "", "", set(), set()),
    ("cadrille", "R_pi + D_t", "0.18", "87.1", "2.1", "0.18", "86.1", "1.5", "0.20", "82.1", "1.4",
     {2, 3, 4, 5, 6, 7, 8, 9, 10}, {2, 3, 4, 5, 6, 7, 8, 9, 10}),
]

# 本复现的两行（黄色高亮）
OURS = [
    ("cadrille (reproduced)", "R_pi + D_t", "0.180", "87.3", "1.86", "0.184", "86.6", "1.44",
     "0.202", "82.3", "2.56"),
    ("CAD-Recode (reproduced)", "R_p", "0.177", "87.9", "2.75", "—", "—", "—", "—", "—", "—"),
]

HEADER1 = ["Method", "Train Data", "Point Cloud", "", "", "Multi-view Images", "", "", "Text", "", ""]
HEADER2 = ["", "", "CD↓", "IoU↑", "IR↓", "CD↓", "IoU↑", "IR↓", "CD↓", "IoU↑", "IR↓"]

YELLOW = "FFF2A8"


def set_font(run, name, size=9, bold=False, underline=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(attr), name)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear",
                                         qn("w:color"): "auto",
                                         qn("w:fill"): hexcolor})
    tcPr.append(shd)


def fill_cell(cell, text, size=8.5, bold=False, underline=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    set_font(r, BODY, size, bold=bold, underline=underline)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    w, h = sec.page_width, sec.page_height
    sec.page_width, sec.page_height = h, w
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.4)

    style = doc.styles["Normal"]
    style.font.name = BODY
    style.font.size = Pt(9)

    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Table 1 · 复现结果与论文对比")
    set_font(r, CN, 15, bold=True, color=(0x1F, 0x33, 0x64))
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Results on DeepCAD test set　|　黄色底纹 = 本复现结果")
    set_font(r, BODY, 9, color=(0x55, 0x55, 0x55))
    p.paragraph_format.space_after = Pt(6)

    # ===== 主表 =====
    n_rows = 2 + len(PAPER) + len(OURS)
    table = doc.add_table(rows=n_rows, cols=11)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 列宽（横向 A4，可用宽度约 26.9 cm）
    col_w = [Cm(4.4), Cm(1.9)] + [Cm(2.28)] * 9
    # 强制固定布局，否则 LibreOffice/Word 会按内容自动调整
    tblPr = table._tbl.tblPr
    layout = tblPr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"})
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = tblPr.makeelement(qn("w:tblW"), {})
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(26.82 * 567)))  # cm -> twips
    # 网格列宽
    grid = table._tbl.tblGrid
    for gc, cw in zip(grid.findall(qn("w:gridCol")), col_w):
        gc.set(qn("w:w"), str(int(cw.cm * 567)))
    for row in table.rows:
        for j, cw in enumerate(col_w):
            row.cells[j].width = cw

    # 两行表头
    for j in range(11):
        fill_cell(table.rows[0].cells[j], HEADER1[j], size=8.5, bold=True)
        fill_cell(table.rows[1].cells[j], HEADER2[j], size=8.5, bold=True)
    # 合并三模态大标题
    for c0, c1 in [(2, 4), (5, 7), (8, 10)]:
        table.rows[0].cells[c0].merge(table.rows[0].cells[c1])
    for j in range(11):
        fill_cell(table.rows[0].cells[j], HEADER1[j] if j < 2 else
                  ("Point Cloud" if j <= 4 else "Multi-view Images" if j <= 7 else "Text"),
                  size=8.5, bold=True)

    # 论文行
    for i, row in enumerate(PAPER):
        cells = table.rows[2 + i].cells
        name, data = row[0], row[1]
        vals = list(row[2:11])
        bold_cols, ul_cols = row[11], row[12]
        for j in range(11):
            if j == 0:
                fill_cell(cells[j], name, align=WD_ALIGN_PARAGRAPH.LEFT)
            elif j == 1:
                fill_cell(cells[j], data)
            else:
                col_idx = j
                fill_cell(cells[j], vals[j - 2],
                          bold=(col_idx in bold_cols),
                          underline=(col_idx in ul_cols))

    # 本复现行（黄色）
    start = 2 + len(PAPER)
    for i, row in enumerate(OURS):
        row_cells = table.rows[start + i].cells
        for j in range(11):
            cell = row_cells[j]
            cell.width = col_w[j]
            shade(cell, YELLOW)
            if j == 0:
                fill_cell(cell, row[0], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            elif j == 1:
                fill_cell(cell, row[1], bold=True)
            else:
                fill_cell(cell, row[j], bold=True)

    # ===== 表注 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Table 1: ")
    set_font(r, BODY, 8.5, bold=True)
    r = p.add_run("Results on DeepCAD test set. The best results are bold, the second best are underlined. "
                  "黄色底纹两行为本复现实测值（10 秒评估超时，无 RL 微调、无 test-time sampling，与论文 Table 1 口径一致）。")
    set_font(r, BODY, 8.5)

    # ===== 第二张表：逐项对比 =====
    doc.add_page_break()
    p = doc.add_paragraph()
    r = p.add_run("本复现结果逐项对比")
    set_font(r, CN, 11, bold=True, color=(0x1F, 0x33, 0x64))
    p.paragraph_format.space_after = Pt(4)

    t2 = doc.add_table(rows=1, cols=9)
    t2.style = "Table Grid"
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    heads = ["行", "模态", "CD↓", "IoU↑", "IR↓", "论文 CD/IoU/IR", "ΔCD", "ΔIoU", "ΔIR"]
    for j, hh in enumerate(heads):
        fill_cell(t2.rows[0].cells[j], hh, size=8.5, bold=True)

    rows2 = [
        ("cadrille R_pi + D_t", "点云", "0.180", "87.3", "1.86", "0.18 / 87.1 / 2.1", "0.000", "+0.2", "−0.24"),
        ("cadrille R_pi + D_t", "图像", "0.184", "86.6", "1.44", "0.18 / 86.1 / 1.5", "+0.004", "+0.5", "−0.06"),
        ("cadrille R_pi + D_t", "文本", "0.202", "82.3", "2.56", "0.20 / 82.1 / 1.4", "+0.002", "+0.2", "+1.16"),
        ("CAD-Recode R_p", "点云", "0.177", "87.9", "2.75", "0.18 / 87.1 / 3.1", "−0.003", "+0.8", "−0.35"),
    ]
    for row in rows2:
        cells = t2.add_row().cells
        for j, v in enumerate(row):
            fill_cell(cells[j], v, size=8.5,
                      align=WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)

    # ===== 说明 =====
    notes = [
        "统计：12 个指标中 11 个与论文一致（91.7%）；CAD-Recode 行三项全部优于论文。",
        "CD 单位已乘 10³；IoU、IR 单位为 %。评估集：DeepCAD test 8046 个样本。",
        "评估超时：本复现使用 10 秒（官方默认 3 秒会导致 IR 虚高约 0.5 个百分点）。",
        "标注说明：官方发布权重对应 R_pi + D_t 配置（三个模态均可推理）。该行与 R_pi 行在点云/图像列数值完全相同，故同时覆盖 R_pi 行的 6 个指标。",
        "唯一偏差：cadrille 文本 IR（2.56 vs 1.4），经 30 秒超时消融确认非超时所致。",
    ]
    for n in notes:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("· " + n)
        set_font(r, BODY, 8.5, color=(0x44, 0x44, 0x44))

    doc.save(DST)
    print("saved:", DST)


if __name__ == "__main__":
    main()
